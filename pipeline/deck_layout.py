"""Slide primitives for the results deck (python-pptx, 16:9)."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

W, H = 13.333, 7.5
MARGIN = 0.6

NAVY = RGBColor(0x10, 0x1B, 0x2E)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK_2 = RGBColor(0x52, 0x51, 0x4E)
INK_3 = RGBColor(0x8A, 0x89, 0x83)
ON_DARK = RGBColor(0xFF, 0xFF, 0xFF)
ON_DARK_2 = RGBColor(0xB9, 0xC6, 0xDA)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
RED = RGBColor(0xE3, 0x49, 0x48)
CARD = RGBColor(0xF1, 0xF5, 0xFB)
SURFACE = RGBColor(0xFC, 0xFC, 0xFB)
FONT = "Calibri"


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def save(self, path: Path) -> None:
        self.prs.save(str(path))

    # -- low level ---------------------------------------------------------------
    def _slide(self, dark: bool = False):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = NAVY if dark else SURFACE
        self.n += 1
        return s

    @staticmethod
    def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing_after=6, bullets=False):
        """`runs` is a string or a list of paragraphs; each paragraph a string or (text, overrides)."""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        paras = [runs] if isinstance(runs, (str, tuple)) else runs
        for i, p in enumerate(paras):
            txt, over = (p, {}) if isinstance(p, str) else p
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para.space_after = Pt(over.get("space_after", spacing_after))
            r = para.add_run()
            r.text = ("•  " if bullets and not over.get("no_bullet") else "") + txt
            f = r.font
            f.name = FONT
            f.size = Pt(over.get("size", size))
            f.bold = over.get("bold", bold)
            f.color.rgb = over.get("color", color)
        return tb

    @staticmethod
    def card(slide, x, y, w, h, fill=CARD):
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.adjustments[0] = 0.06
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    @staticmethod
    def image(slide, path: Path, x, y, max_w, max_h, align="center"):
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(max_w / iw, max_h / ih)
        w, h = iw * scale, ih * scale
        dx = (max_w - w) / 2 if align == "center" else 0
        dy = (max_h - h) / 2
        return slide.shapes.add_picture(str(path), Inches(x + dx), Inches(y + dy), Inches(w), Inches(h))

    def _footer(self, slide, source: str | None, dark=False):
        if source:
            self.text(slide, MARGIN, H - 0.42, W - 2 * MARGIN - 0.6, 0.3, source, size=9,
                      color=ON_DARK_2 if dark else INK_3)
        self.text(slide, W - MARGIN - 0.5, H - 0.42, 0.5, 0.3, str(self.n), size=9,
                  color=ON_DARK_2 if dark else INK_3, align=PP_ALIGN.RIGHT)

    # -- slide types --------------------------------------------------------------
    def title_slide(self, title, subtitle, byline):
        s = self._slide(dark=True)
        self.text(s, MARGIN + 0.2, 2.0, W - 2 * MARGIN - 1.5, 2.2, title, size=40, bold=True, color=ON_DARK,
                  anchor=MSO_ANCHOR.BOTTOM)
        self.text(s, MARGIN + 0.2, 4.45, W - 2 * MARGIN - 2.5, 1.0, subtitle, size=18, color=ON_DARK_2)
        self.text(s, MARGIN + 0.2, H - 1.1, 8, 0.4, byline, size=13, color=ON_DARK_2)
        return s

    def header(self, s, title, subtitle=None, dark=False):
        self.text(s, MARGIN, 0.45, W - 2 * MARGIN, 0.95, title, size=28, bold=True,
                  color=ON_DARK if dark else INK, anchor=MSO_ANCHOR.TOP)
        if subtitle:
            self.text(s, MARGIN, 1.32, W - 2 * MARGIN, 0.5, subtitle, size=15, color=ON_DARK_2 if dark else INK_2)

    def figure_slide(self, title, fig: Path, subtitle=None, notes: list[str] | None = None, source=None):
        """Figure left, optional takeaway card right."""
        s = self._slide()
        self.header(s, title, subtitle)
        top = 1.95 if subtitle else 1.6
        if notes:
            fig_w = 8.3
            self.image(s, fig, MARGIN, top, fig_w, H - top - 0.65)
            cx = MARGIN + fig_w + 0.35
            cw = W - MARGIN - cx
            chars = sum(len(n) for n in notes)
            ch = min(H - top - 0.75, 0.7 + 0.42 * len(notes) + chars / 34 * 0.27)
            self.card(s, cx, top, cw, ch)
            self.text(s, cx + 0.3, top + 0.3, cw - 0.6, ch - 0.5, notes, size=14, color=INK, spacing_after=12,
                      bullets=True)
        else:
            self.image(s, fig, MARGIN, top, W - 2 * MARGIN, H - top - 0.65)
        self._footer(s, source)
        return s

    def stats_slide(self, title, stats: list[tuple[str, str, str]], subtitle=None, source=None, dark=False):
        """stats: (big number, label, detail)."""
        s = self._slide(dark=dark)
        self.header(s, title, subtitle, dark=dark)
        n = len(stats)
        gap = 0.35
        cw = (W - 2 * MARGIN - gap * (n - 1)) / n
        top = 2.3 if subtitle else 2.0
        for i, (big, label, detail) in enumerate(stats):
            x = MARGIN + i * (cw + gap)
            self.card(s, x, top, cw, 3.0, fill=RGBColor(0x1B, 0x2A, 0x44) if dark else CARD)
            self.text(s, x + 0.3, top + 0.35, cw - 0.6, 1.1, big, size=40, bold=True,
                      color=ON_DARK if dark else BLUE)
            self.text(s, x + 0.3, top + 1.5, cw - 0.6, 0.8, label, size=16, bold=True,
                      color=ON_DARK if dark else INK)
            self.text(s, x + 0.3, top + 2.1, cw - 0.6, 0.85, detail, size=13, color=ON_DARK_2 if dark else INK_2)
        self._footer(s, source, dark=dark)
        return s

    def columns_slide(self, title, cols: list[tuple[str, list[str]]], subtitle=None, source=None, dark=False):
        s = self._slide(dark=dark)
        self.header(s, title, subtitle, dark=dark)
        n = len(cols)
        gap = 0.35
        cw = (W - 2 * MARGIN - gap * (n - 1)) / n
        top = 2.1 if subtitle else 1.8
        for i, (head, items) in enumerate(cols):
            x = MARGIN + i * (cw + gap)
            self.card(s, x, top, cw, H - top - 0.8, fill=RGBColor(0x1B, 0x2A, 0x44) if dark else CARD)
            self.text(s, x + 0.3, top + 0.3, cw - 0.6, 0.6, head, size=18, bold=True,
                      color=ON_DARK if dark else INK)
            self.text(s, x + 0.3, top + 1.0, cw - 0.6, H - top - 2.0, items, size=16 if n <= 2 else 15,
                      color=ON_DARK_2 if dark else INK_2, spacing_after=14, bullets=True)
        self._footer(s, source, dark=dark)
        return s

    def flow_slide(self, title, steps: list[tuple[str, str]], subtitle=None, fig: Path | None = None, source=None):
        """Horizontal numbered process with optional figure below."""
        s = self._slide()
        self.header(s, title, subtitle)
        n = len(steps)
        gap = 0.2
        top = 1.95 if subtitle else 1.65
        cw = (W - 2 * MARGIN - gap * (n - 1)) / n
        ch = 1.75
        for i, (head, body) in enumerate(steps):
            x = MARGIN + i * (cw + gap)
            self.card(s, x, top, cw, ch)
            self.text(s, x + 0.18, top + 0.15, cw - 0.36, 0.35, f"{i + 1:02d}", size=12, bold=True, color=BLUE)
            self.text(s, x + 0.18, top + 0.45, cw - 0.36, 0.45, head, size=13.5, bold=True, color=INK)
            self.text(s, x + 0.18, top + 0.9, cw - 0.36, ch - 0.95, body, size=10.5, color=INK_2)
        if fig is not None:
            self.image(s, fig, MARGIN, top + ch + 0.3, W - 2 * MARGIN, H - (top + ch + 0.3) - 0.6)
        self._footer(s, source)
        return s
